"""Tarea 268 — tres menores de pantalla y reportes.

[F-2] Paper pintaba un precio faltante como P&L ``+$0.00 / +0.00%`` en verde. [F-3]
``reports/dashboard_sim_principal.html`` era una foto de mayo de la cuenta cerrada. Y, de la 277,
el Excel armaba DOS hojas de transacciones —una siempre, aunque ``tx_history`` estuviera apagado—
con totales que ponían mal el signo de la comisión en una de las dos patas.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest

from database.models import Portfolio, Position, Transaction, session_scope
from reports.excel_report import total_de_transaccion
from ui.paper_tab import textos_pnl

_REPO = Path(__file__).resolve().parent.parent


def test_sin_precio_el_PNL_no_es_un_cero_verde():
    assert textos_pnl(None, 10, 50.0) == ("—", "—", None)


def test_con_precio_el_PNL_tiene_signo_y_color():
    assert textos_pnl(60.0, 10, 50.0) == ("+$100.00", "+20.00%", True)
    assert textos_pnl(40.0, 10, 50.0) == ("-$100.00", "-20.00%", False)


def test_el_total_suma_la_comision_en_la_compra_y_la_resta_en_la_venta():
    assert total_de_transaccion("BUY", 10, 100.0, 1.0) == pytest.approx(1001.0)
    assert total_de_transaccion("SELL", 4, 120.0, 1.0) == pytest.approx(479.0)
    assert total_de_transaccion("SELL", 4, 120.0, None) == pytest.approx(480.0)


def _cartera():
    with session_scope() as s:
        pf = Portfolio(name="Mis Acciones")
        s.add(pf)
        s.flush()
        pos = Position(portfolio_id=pf.id, ticker="AAA", quantity=6, avg_buy_price=100.0)
        s.add(pos)
        s.flush()
        s.add(
            Transaction(
                position_id=pos.id,
                transaction_type="BUY",
                quantity=10,
                price=100.0,
                fees=1.0,
                date=datetime(2026, 9, 1),
            )
        )
        s.add(
            Transaction(
                position_id=pos.id,
                transaction_type="SELL",
                quantity=4,
                price=120.0,
                fees=1.0,
                date=datetime(2026, 9, 5),
            )
        )
    with session_scope() as s:
        out = s.query(Position).all()
        s.expunge_all()
        return out


def _hojas(path):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    try:
        return {ws.title: [[c.value for c in r] for r in ws.iter_rows()] for ws in wb.worksheets}
    finally:
        wb.close()


@pytest.mark.parametrize("incluir,esperadas", [(True, 1), (False, 0)])
def test_el_excel_tiene_UNA_hoja_de_transacciones_y_sólo_si_se_pide(test_db, tmp_path, incluir, esperadas):
    from reports.excel_report import generate_portfolio_excel

    out = tmp_path / "r.xlsx"
    generate_portfolio_excel(str(out), "Mis Acciones", _cartera(), {}, "USD", include_tx=incluir)
    hojas = _hojas(out)
    assert sum(1 for t in hojas if t.startswith("Transacciones")) == esperadas


def test_la_hoja_usa_el_total_con_el_signo_correcto(test_db, tmp_path):
    from reports.excel_report import generate_portfolio_excel

    out = tmp_path / "r.xlsx"
    generate_portfolio_excel(str(out), "Mis Acciones", _cartera(), {}, "USD", include_tx=True)
    filas = _hojas(out)["Transacciones"]
    totales = {f[1]: f[5] for f in filas[1:] if f[1] in ("BUY", "SELL")}
    assert totales == {"BUY": pytest.approx(1001.0), "SELL": pytest.approx(479.0)}


def test_el_dashboard_viejo_de_la_cuenta_cerrada_ya_no_esta():
    assert not (_REPO / "reports" / "dashboard_sim_principal.html").exists()


def test_el_resumen_del_reporte_no_valua_al_costo_un_precio_faltante(test_db, tmp_path):
    """Hallado al cerrar la 268: el resumen de Excel y PDF tenía el defecto de la 281 [P-1]."""
    from database.cartera_real import valor_y_pl
    from reports.excel_report import generate_portfolio_excel

    posiciones = _cartera()  # AAA: 6 acciones a 100
    sin = valor_y_pl(posiciones, {})
    assert sin["valor"] == 0.0 and sin["sin_precio"] == ["AAA"], "sin precio entró al valor al costo"
    con = valor_y_pl(posiciones, {"AAA": {"price": 110.0}})
    assert con["valor"] == pytest.approx(660.0) and con["pl_pct"] == pytest.approx(10.0)

    out = tmp_path / "r.xlsx"
    generate_portfolio_excel(str(out), "Mis Acciones", posiciones, {}, "USD", include_tx=False)
    texto = " ".join(str(v) for filas in _hojas(out).values() for f in filas for v in f if v is not None)
    assert "1 sin precio: AAA" in texto
